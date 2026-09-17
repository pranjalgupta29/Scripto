"""PDF and user-pasted adapters."""

from __future__ import annotations

import io

import httpx

from scripto.adapters.base import FetchError, ParsedSource, SourceAdapter
from scripto.config import settings


class PdfAdapter(SourceAdapter):
    type = "pdf"

    def fetch(self, url: str) -> bytes:
        try:
            response = httpx.get(
                url,
                headers={"User-Agent": settings.http_user_agent},
                timeout=settings.http_timeout_seconds,
                follow_redirects=True,
            )
            response.raise_for_status()
        except httpx.HTTPError as exc:
            raise FetchError(f"fetch failed: {exc}") from exc
        return response.content

    def parse(self, raw: bytes, *, url: str | None = None) -> ParsedSource:
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(raw))
            pages = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:
            raise FetchError(f"pdf parse failed: {exc}") from exc

        text = "\n\n".join(p.strip() for p in pages if p.strip())
        if not text:
            raise FetchError("pdf contained no extractable text")

        meta = reader.metadata or {}
        return ParsedSource(
            title=str(meta.get("/Title")) if meta.get("/Title") else None,
            author=str(meta.get("/Author")) if meta.get("/Author") else None,
            published_at=None,
            text=text,
            segments=None,
        )


class DocxAdapter(SourceAdapter):
    """A Word document the host uploaded: a resume, a bio, a briefing note.

    Like pasted text, it arrives with its content, so there is nothing to fetch.
    """

    type = "docx"

    def fetch(self, url: str) -> bytes:
        raise FetchError("docx sources are uploaded with their content, never fetched")

    def parse(self, raw: bytes, *, url: str | None = None) -> ParsedSource:
        import docx

        try:
            document = docx.Document(io.BytesIO(raw))
        except Exception as exc:
            raise FetchError(f"docx parse failed: {exc}") from exc

        blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        # Resumes keep dates, employers and skills in tables, which the
        # paragraph list alone would miss entirely.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))

        text = "\n\n".join(blocks)
        if not text:
            raise FetchError("docx contained no extractable text")

        core = document.core_properties
        return ParsedSource(
            title=(core.title or None) or (blocks[0][:200] if blocks else None),
            author=core.author or None,
            published_at=None,
            text=text,
            segments=None,
        )


class UserPastedAdapter(SourceAdapter):
    """Text the user pasted directly. Already in hand, so fetch is a no-op.

    This is also the path the thin-footprint mode leans on: a pasted bio, CV or
    internal note becomes a first-class source with no special casing.
    """

    type = "user_pasted"

    def fetch(self, url: str) -> bytes:
        raise FetchError("user_pasted sources are created with their content, never fetched")

    def parse(self, raw: bytes, *, url: str | None = None) -> ParsedSource:
        text = raw.decode("utf-8", errors="replace").strip()
        if not text:
            raise FetchError("pasted text was empty")
        first_line = text.splitlines()[0].strip()
        return ParsedSource(
            title=first_line[:200] if first_line else None,
            author=None,
            published_at=None,
            text=text,
            segments=None,
        )
